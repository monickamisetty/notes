#!/usr/bin/env python3
"""
CCV Engineering Assistant Bootstrap

Creates a generic VS Code / GitHub Copilot workspace for CCV
(Cross-Coverage Validation) development and validation.

Contains NO internal rules, requirements, Guidewire data, URLs, credentials,
or company/client-specific source code.

Usage:
    python bootstrap_ccv_assistant.py
    python bootstrap_ccv_assistant.py --root ccv-assistant
    python bootstrap_ccv_assistant.py --force

Optional Excel dependency:
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import textwrap
from pathlib import Path
from typing import Dict

FILES: Dict[str, str] = {}

FILES['README.md'] = r'''
# CCV Engineering Assistant

Repository-based CCV engineering assistant for VS Code / GitHub Copilot.

## Core capabilities

- business requirement -> candidate CCV JSON Logic
- rule explanation and semantic validation
- custom JSON Logic operation awareness
- `x-gw-*` CCV response-contract awareness
- runtime test-scenario generation
- requirement-to-rule and rule-to-requirement analysis
- React/UI-coded CCV behavior knowledge
- coverage/state validation
- cross-state comparison
- known-defect pattern checking

## State maturity

| State | Status |
|---|---|
| WI | RELEASED / reference |
| IL | DEVELOPMENT |
| OH | DEVELOPMENT |
| AZ | DEVELOPMENT |
| UT | REFINEMENT |
| IN | REFINEMENT |

Wisconsin is an established implementation reference. A difference from
Wisconsin is not automatically a defect; target-state requirements remain the
authority for intended behavior.

## Design model

```text
Business Requirement
        |
        v
Expected CCV Behavior
        |
   +----+----+
   |         |
   v         v
JSON Logic  React/UI
   |         |
   +----+----+
        |
        v
Runtime behavior + tests
```

## First setup inside work

1. Put Postman/system-table exports in `states/<STATE>/raw/ccv-rules.json`.
2. Put Guidewire/coverage exports in `states/<STATE>/raw/coverages.json`.
3. Fill `knowledge/json-logic/custom-operators.yaml`.
4. Verify `knowledge/ccv-response/response-contract.yaml`.
5. Add UI-coded scenarios to `knowledge/implementation-patterns/ui-coded-rules.md`.
6. Import the existing internal Excel requirements with `tools/import_requirements.py`.
7. Run `python tools/build_indexes.py`.
8. Validate representative Wisconsin behavior before using the assistant on development states.

## Requirement import

Inspect workbook sheets/headers:

```bash
python tools/import_requirements.py path/to/requirements.xlsx --inspect
```

Import with explicit columns:

```bash
python tools/import_requirements.py path/to/requirements.xlsx \
  --sheet "Requirements" \
  --state-column "State" \
  --coverage-column "Coverage" \
  --id-column "Requirement ID" \
  --text-column "Requirement"
```

If one sheet is state-specific:

```bash
python tools/import_requirements.py path/to/requirements.xlsx \
  --sheet "Illinois" \
  --state IL \
  --coverage-column "Coverage" \
  --id-column "Requirement ID" \
  --text-column "Requirement"
```

## Example Copilot requests

- Explain Wisconsin UM rules. Retrieve the relevant files first.
- Validate IL UM against its requirements.
- Convert this business requirement into candidate CCV JSON Logic.
- Review this rule for AND/OR, boundary, null, default, and forbidden issues.
- Generate runtime tests for this rule.
- Find the closest released Wisconsin implementation pattern.
- Compare IL and WI for this specific behavior.
'''

FILES['.gitignore'] = r'''
__pycache__/
*.pyc
.venv/
venv/
.env
.DS_Store
reports/*.json
reports/*.md
'''

FILES['.github/copilot-instructions.md'] = r'''
# CCV Engineering Assistant — Global Instructions

You are working in a CCV (Cross-Coverage Validation) engineering repository.

## Source-of-truth rules

1. Never invent existing CCV behavior.
2. Retrieve relevant repository files before answering about existing behavior.
3. Treat state-specific business requirements as the authority for intended behavior.
4. Treat implementation as evidence of how behavior was built, not proof the requirement is correct.
5. CCV behavior can be implemented in JSON Logic, React/UI code, or both.
6. Check documented UI-coded behavior before claiming that a requirement has no implementation.

## State maturity

- WI: RELEASED; may be used as an implementation reference.
- IL: DEVELOPMENT.
- OH: DEVELOPMENT.
- AZ: DEVELOPMENT.
- UT: REFINEMENT.
- IN: REFINEMENT.

A difference from Wisconsin is a review signal, not proof of a defect.

## JSON Logic

Before rejecting an operator as unsupported:

1. Check `knowledge/json-logic/custom-operators.yaml`.
2. If documented there, validate using the documented CCV semantics.
3. Do not infer undocumented custom-operator behavior.

## CCV response/property contract

Use `knowledge/ccv-response/response-contract.yaml`.

Known names:

- `x-gw-cc-rule-info`
- `x-gw-cc-rule-success`
- `x-gw-cc-rule-error`
- `x-gw-default`
- `x-gw-forbidden`
- `x-gw-forbidden-property`

Do not invent undocumented meanings or property values.

## Validation dimensions

When reviewing a rule, check:

1. JSON structure
2. standard/custom operation usage
3. requirement-to-rule semantic match
4. data-model/property references
5. condition behavior and boundaries
6. CCV response/effect correctness
7. cross-rule interactions
8. React/UI interactions
9. runtime-test coverage

Pay special attention to:

- AND vs OR
- > vs >=
- < vs <=
- equality boundaries
- missing/null values
- decline behavior
- filtering an option versus hiding an entire coverage
- filtering options without correcting an invalid existing selection
- x-gw-default behavior
- x-gw-forbidden / x-gw-forbidden-property behavior
- duplicate, overlapping, or conflicting rules
- UI and JSON Logic both changing the same behavior
- display value versus underlying Guidewire/system value

## Verdicts

Use exactly one primary verdict:

- PASS
- CONCERN
- LIKELY DEFECT
- CANNOT VERIFY

PASS means no obvious issue was found from available evidence. It does not prove runtime correctness.

## Runtime verification

For rule generation or validation, create relevant tests from these categories:

- happy path
- negative cases
- each independent condition false
- equality boundary
- one above / one below boundary
- missing field
- null
- decline
- invalid pre-existing dependent selection
- custom-operation-specific cases

Recommend running generated cases through the actual CCV/JSON Logic runtime.

## Retrieval efficiency

Use index files before reading large raw exports. Read the smallest useful set of files.

Prefer this validation answer structure:

1. Verdict
2. Requirement interpretation
3. Implementation found
4. Findings
5. Reference patterns
6. Recommended correction
7. Runtime tests
'''

FILES['.github/agents/ccv-expert.agent.md'] = r'''
---
description: Specialized CCV agent for requirements, JSON Logic, UI-coded behavior, rule/state validation, testing, and comparison.
---

# CCV Expert

You are the specialized CCV engineering agent for this repository.

## Responsibilities

- explain rules
- interpret requirements
- generate candidate CCV JSON Logic
- validate rules against requirements
- understand documented custom operations
- understand documented `x-gw-*` response behavior
- recognize React/UI-coded CCV behavior
- generate runtime tests
- map requirements to implementations
- find duplicates/overlaps/conflicts
- compare states
- validate a coverage
- validate a state

## Existing behavior workflow

1. Retrieve indexes first.
2. Retrieve the exact requirement(s).
3. Retrieve relevant rule(s).
4. Retrieve custom-operator definitions if used.
5. Retrieve response-contract definitions if `x-gw-*` properties are involved.
6. Retrieve relevant UI-coded behavior.
7. Retrieve Wisconsin references only where useful.
8. Compare evidence and produce a bounded verdict.

## Requirement -> rule workflow

1. Normalize trigger, conditions, target, effect, and boundaries.
2. Search for similar existing behavior.
3. Check custom operators.
4. Determine whether JSON Logic, UI code, or a combination is the better fit.
5. Generate candidate logic only where supported by available field/property knowledge.
6. Explain each significant condition.
7. List assumptions and concerns.
8. Generate runtime tests.

Never silently invent a field mapping.

## Validation verdict

Use one of:

`PASS`, `CONCERN`, `LIKELY DEFECT`, `CANNOT VERIFY`.

Wisconsin is a released implementation reference, not a universal state-rule authority.
'''

FILES['.github/skills/create-rule/SKILL.md'] = r'''
# Create CCV Rule

Use when converting a textual business requirement into a candidate CCV implementation.

## Procedure

1. Retrieve requirement and relevant coverage data.
2. Normalize:
   - trigger
   - conditions
   - target
   - expected effect
   - boundary behavior
   - missing/null behavior if specified
3. Search similar rules and UI-coded patterns.
4. Check custom JSON Logic operations.
5. Decide whether JSON Logic is a good fit.
6. Do not force dynamic dependent-option/reselection behavior into JSON Logic if a documented UI pattern is more appropriate.
7. Generate candidate rule.
8. Explain requirement-to-condition mapping.
9. List assumptions/concerns.
10. Generate runtime tests.

## Output

- READY FOR RUNTIME TEST or NEEDS REVIEW
- requirement interpretation
- candidate JSON Logic
- expected `x-gw-*` effect
- similar reference
- assumptions
- concerns
- tests
'''

FILES['.github/skills/validate-rule/SKILL.md'] = r'''
# Validate CCV Rule

Compare:

`Requirement -> Expected Behavior -> Actual Implementation`

Implementation may be JSON Logic, React/UI, or both.

## Check

- JSON structure
- documented standard/custom operators
- AND/OR semantics
- >/>= and </<= boundaries
- inclusion/exclusion
- required/optional/decline
- x-gw-default behavior
- x-gw-forbidden behavior and target/property
- missing requirement branches
- inverted conditions
- coverage/property/value references
- null/missing behavior
- duplicate/overlap/conflict
- UI dynamic filtering
- invalid selected value after filtering
- JSON/UI double handling

## Verdict

Use exactly one:

- PASS
- CONCERN
- LIKELY DEFECT
- CANNOT VERIFY

Always include targeted runtime tests.
'''

FILES['.github/skills/generate-tests/SKILL.md'] = r'''
# Generate CCV Runtime Tests

Generate only relevant categories, but consider:

1. happy path
2. negative path
3. each condition independently false
4. equality boundary
5. just below boundary
6. just above boundary
7. missing property
8. null property
9. empty value/list
10. decline
11. dependent value already valid
12. dependent value becomes invalid
13. custom-operation-specific edge cases
14. UI option-filter/reselection cases

For each test include:

- name
- purpose
- input data
- expected rule result
- expected `x-gw-*` effect if known
- notes

Prefer JSON payloads copyable into the real validation tool.
Mark unverifiable expected behavior as `CANNOT VERIFY`.
'''

FILES['.github/skills/analyze-rules/SKILL.md'] = r'''
# Analyze CCV Rules

Use for:

- explain a rule
- list rules by state/coverage
- find rules using an operator
- find rules using a specific `x-gw-*` property
- find similar rules
- identify dependencies
- identify UI-coded behavior
- identify overlaps/conflicts
- identify requirements without obvious implementation
- identify rules without mapped requirements

## Retrieval

1. Read indexes first.
2. Retrieve only candidate rules.
3. Retrieve relevant requirements.
4. Retrieve custom operators and response-contract knowledge as needed.
5. Retrieve UI-coded patterns.
6. Retrieve Wisconsin references only when useful.

Always identify source rule IDs/files.
'''

FILES['.github/skills/validate-state/SKILL.md'] = r'''
# Validate CCV State

Use for development/refinement state validation.

Prefer coverage-level validation first when the state is large.

## Analyze

- requirements
- coverage model
- CCV rules
- custom operators
- response contract
- UI-coded behavior
- known issues
- relevant Wisconsin references

## Identify

- requirement with no obvious implementation
- rule with no mapped requirement
- malformed/suspicious rules
- undocumented operators
- questionable property/value references
- missing branches
- default-value omissions
- forbidden-target mistakes
- duplicates/overlaps/conflicts
- JSON/UI double handling
- missing UI-coded implementation knowledge
- meaningful differences from Wisconsin references
- runtime test gaps

A Wisconsin difference is a review signal, not proof of a defect.
'''

FILES['.github/skills/compare-states/SKILL.md'] = r'''
# Compare CCV States

Compare:

- requirements
- coverages
- business condition
- rule structure
- custom operations
- `x-gw-*` effect
- default behavior
- forbidden behavior
- React/UI behavior
- tests

Use each state's own requirement as authority.
Classify differences as:

- expected/state-specific
- implementation-pattern difference
- concern
- likely defect
- cannot verify
'''

FILES['knowledge/ccv/states.yaml'] = r'''
states:
  WI:
    name: Wisconsin
    status: RELEASED
    reference: true
  IL:
    name: Illinois
    status: DEVELOPMENT
    reference: false
  OH:
    name: Ohio
    status: DEVELOPMENT
    reference: false
  AZ:
    name: Arizona
    status: DEVELOPMENT
    reference: false
  UT:
    name: Utah
    status: REFINEMENT
    reference: false
  IN:
    name: Indiana
    status: REFINEMENT
    reference: false
'''

FILES['knowledge/ccv/architecture.md'] = r'''
# CCV Architecture

Complete this with verified internal architecture.

Capture:

- where rules originate
- system-table storage model
- Guidewire inputs
- how rules are retrieved/evaluated
- where custom operations run
- how the UI consumes rule results
- when React/UI implements business behavior directly
- how runtime validation is performed

## Conceptual model

```text
Requirement
   |
Expected behavior
   |
 +---+---+
 |       |
JSON    React/UI
Logic
 |       |
 +---+---+
     |
User/runtime behavior
```
'''

FILES['knowledge/ccv/evaluation-flow.md'] = r'''
# CCV Evaluation Flow

Document the verified runtime sequence.

Answer internally:

1. What data is passed to JSON Logic?
2. Where do custom operations execute?
3. What is the exact rule result shape?
4. How are multiple matching rules combined?
5. How does UI consume each `x-gw-*` property?
6. What happens when multiple rules affect one target?
7. When does React/UI-coded behavior run?
8. How does initial load differ from user selection changes?
'''

FILES['knowledge/json-logic/custom-operators.yaml'] = r'''
# Add EVERY custom JSON Logic operator implemented by the CCV engine.
# Do not infer undocumented behavior.
#
# Example template:
# operators:
#   - name: exampleOperator
#     purpose: >
#       Exact purpose.
#     syntax:
#       example:
#         exampleOperator:
#           - arg1
#           - arg2
#     arguments:
#       - name: arg1
#         type: string
#         required: true
#     returnType: boolean
#     nullBehavior: "DOCUMENT"
#     missingValueBehavior: "DOCUMENT"
#     examples: []
#     edgeCases: []
#     implementationReference: "OPTIONAL INTERNAL SOURCE PATH"

operators: []
'''

FILES['knowledge/ccv-response/response-contract.yaml'] = r'''
# Verify and enrich these properties from the actual internal implementation.

properties:
  x-gw-cc-rule-success:
    category: availability
    verified: false
    description: >
      Used by CCV/UI availability behavior. Document exact value structure and
      presence/absence behavior.
    examples: []

  x-gw-cc-rule-info:
    category: informational-message
    verified: false
    description: >
      Used for informational UI behavior. Document exact code/message mapping.
    examples: []

  x-gw-cc-rule-error:
    category: validation-message
    verified: false
    description: >
      Used for error/warning UI behavior. Document exact severity/rendering semantics.
    examples: []

  x-gw-default:
    category: value-mutation
    verified: false
    description: >
      Sets a target coverage/property value when a condition matches. Document
      exact target/value representation.
    examples: []

  x-gw-forbidden:
    category: restriction
    verified: false
    description: >
      Identifies a forbidden/restricted target. Document exact runtime shape.
    examples: []

  x-gw-forbidden-property:
    category: restriction-property
    verified: false
    description: >
      Document exact allowed values and relationship to x-gw-forbidden. Do not
      infer coverage/term/choice semantics until verified.
    allowedValues: []
    examples: []
'''

FILES['knowledge/implementation-patterns/ui-coded-rules.md'] = r'''
# React/UI-Coded CCV Behavior

CCV behavior is not limited to JSON Logic. Capture business scenarios implemented
in React/UI or split between JSON Logic and UI.

## Pattern template

### ID
`<STATE>-UI-CCV-###`

### State
`<STATE>`

### Coverages
- `<coverage>`

### Business behavior
Describe behavior independently of code.

### Trigger
What user/data change causes it?

### Allowed values / restrictions
Describe filtering/restrictions.

### Existing-selection behavior
What happens if a current selection becomes invalid?

### Special cases
Boundaries/exceptions.

### Implementation type
`UI` or `JSON_LOGIC_AND_UI`

### Internal source reference
Add React component/hook/utility paths internally.

### Tests
Representative cases.

---

## Wisconsin BI / UM / UIM — VERIFY INTERNALLY

### ID
`WI-UI-BI-UM-UIM`

### Status
Released reference pattern.

### Coverages
- BI — Bodily Injury
- UM — Uninsured Motorist
- UIM — Underinsured Motorist

### UM/UIM maximum controlled by BI

UM and UIM should not expose selectable limit options above the currently
selected BI limit.

Conceptually:

`UM <= BI`

`UIM <= BI`

If BI is reduced and the current UM/UIM selection becomes invalid, the dependent
coverage should move to the highest remaining allowed value, subject to the
verified state-specific rules.

Filtering the dropdown alone may be insufficient if the current value is now invalid.

### UIM relationship to UM

Verify the exact Wisconsin implementation internally.

Starter behavior to verify:

- UM minimum 25/50 -> UIM non-decline value 50/100
- for other applicable UM values, UIM generally corresponds to UM
- decline may also be available according to the actual rule

### Why UI implementation is relevant

Dynamic dependent-option filtering and reselection can require React/UI logic
rather than static JSON Logic alone.

### Internal source reference
Add actual source paths here.

### Required tests

- BI high; UM/UIM same value
- BI reduced below current UM
- BI reduced below current UIM
- dependent value still valid
- dependent value becomes invalid
- UM minimum value
- UIM special minimum case
- decline
- initial load
- repeated increase/decrease
'''

FILES['knowledge/validation/known-issues.md'] = r'''
# CCV Known Defect Patterns

Use as a review checklist, not automatic proof of a defect.

## Boundary operator mismatch
Requirement says "at least" but implementation uses `>` instead of `>=`, etc.

## AND implemented as OR
All business conditions are required, but rule uses `or`.

## Inverted comparison
Controlling/dependent coverage order is reversed.

## Decline restriction at wrong level
Requirement makes coverage mandatory, but implementation hides entire coverage
instead of forbidding/restricting decline.

## Filtered options but stale invalid selection
Options are filtered after a controlling value changes, but the selected dependent
value remains outside the allowed set.

## Duplicate/overlapping rule
Multiple rules trigger on overlapping conditions and produce competing effects.

## JSON Logic and UI double handling
Both layers mutate/restrict the same target with sequencing risk.

## Display value versus system code
Rule compares a label while runtime data uses an underlying Guidewire/system code.

## Missing/null behavior
Rule does not safely account for missing/null data.

## Default without validity check
A default sets a dependent value that may violate another active rule.
'''

FILES['requirements/README.md'] = r'''
# Requirements

Requirements remain inside the authorized work environment.

Use `tools/import_requirements.py` to read the existing Excel workbook and create
small normalized JSON files. The importer preserves original requirement text
and does not use AI to reinterpret semantics.
'''

FILES['templates/requirement.template.json'] = r'''
{
  "id": "STATE-REQ-001",
  "state": "STATE",
  "coverage": "COVERAGE",
  "rawRequirement": "Exact business requirement text",
  "status": "",
  "businessNotes": "",
  "implementationStatus": "UNKNOWN",
  "implementationType": "UNKNOWN",
  "ruleIds": [],
  "uiBehaviorIds": [],
  "source": {
    "workbook": "",
    "sheet": "",
    "row": null
  }
}
'''

FILES['schemas/ccv-rule.schema.json'] = r'''
{
  "$schema": "https://json-schema.org/draft/2020-12/schema",
  "title": "Generic CCV Rule Container",
  "description": "Starter schema. Adapt to the actual internal rule export shape.",
  "type": ["object", "array"]
}
'''

FILES['schemas/test-case.schema.json'] = r'''
{
  "$schema": "https://json-schema.org/draft/2020-12/schema",
  "title": "CCV Runtime Test Case",
  "type": "object",
  "properties": {
    "name": {"type": "string"},
    "purpose": {"type": "string"},
    "input": {},
    "expectedMatch": {},
    "expectedOutput": {},
    "notes": {"type": "string"}
  },
  "required": ["name", "input"]
}
'''

FILES['tools/import_requirements.py'] = r'''#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

ROOT = Path(__file__).resolve().parents[1]


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def safe_id(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_.-]+", "-", value.strip())
    return value.strip("-") or "REQ"


def load_wb(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required. Install with: pip install openpyxl")
    return load_workbook(path, read_only=True, data_only=True)


def inspect(path: Path) -> None:
    wb = load_wb(path)
    print(f"Workbook: {path}")
    for ws in wb.worksheets:
        print(f"\nSheet: {ws.title}")
        shown = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True):
            vals = [clean(v) for v in row]
            if any(vals):
                print("  " + " | ".join(vals))
                shown += 1
                if shown >= 3:
                    break


def find_header_row(ws, required: List[str], max_rows: int = 30) -> int:
    needed = {h.strip().lower() for h in required if h}
    for row_num in range(1, min(ws.max_row, max_rows) + 1):
        present = {clean(c.value).lower() for c in ws[row_num] if clean(c.value)}
        if needed.issubset(present):
            return row_num
    raise ValueError("Could not find headers: " + ", ".join(required))


def header_map(ws, row_num: int) -> Dict[str, int]:
    result = {}
    for idx, cell in enumerate(ws[row_num], start=1):
        name = clean(cell.value)
        if name:
            result[name.lower()] = idx
    return result


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("workbook", type=Path)
    p.add_argument("--inspect", action="store_true")
    p.add_argument("--sheet")
    p.add_argument("--header-row", type=int)
    p.add_argument("--state")
    p.add_argument("--state-column", default="State")
    p.add_argument("--coverage-column", default="Coverage")
    p.add_argument("--id-column", default="Requirement ID")
    p.add_argument("--text-column", default="Requirement")
    p.add_argument("--status-column", default="Status")
    p.add_argument("--notes-column", default="Business Notes")
    p.add_argument("--overwrite", action="store_true")
    args = p.parse_args()

    if not args.workbook.exists():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    if args.inspect:
        inspect(args.workbook)
        return 0

    wb = load_wb(args.workbook)
    sheets = [wb[args.sheet]] if args.sheet else list(wb.worksheets)
    imported = skipped = 0

    for ws in sheets:
        required = [args.coverage_column, args.id_column, args.text_column]
        if not args.state:
            required.append(args.state_column)
        try:
            header_row = args.header_row or find_header_row(ws, required)
        except ValueError as exc:
            print(f"Skipping {ws.title}: {exc}")
            continue

        headers = header_map(ws, header_row)
        col = lambda name: headers.get(name.strip().lower())
        state_col = None if args.state else col(args.state_column)
        coverage_col = col(args.coverage_column)
        id_col = col(args.id_column)
        text_col = col(args.text_column)
        status_col = col(args.status_column)
        notes_col = col(args.notes_column)

        def val(row: int, column: Optional[int]) -> str:
            return "" if not column else clean(ws.cell(row=row, column=column).value)

        for row in range(header_row + 1, ws.max_row + 1):
            raw_text = val(row, text_col)
            req_id = val(row, id_col)
            if not raw_text and not req_id:
                continue

            state = clean(args.state) or val(row, state_col)
            coverage = val(row, coverage_col)
            if not state:
                skipped += 1
                continue
            if not req_id:
                req_id = f"{state}-REQ-{row:04d}"

            data = {
                "id": req_id,
                "state": state,
                "coverage": coverage,
                "rawRequirement": raw_text,
                "status": val(row, status_col),
                "businessNotes": val(row, notes_col),
                "implementationStatus": "UNKNOWN",
                "implementationType": "UNKNOWN",
                "ruleIds": [],
                "uiBehaviorIds": [],
                "source": {
                    "workbook": args.workbook.name,
                    "sheet": ws.title,
                    "row": row
                }
            }

            out_dir = ROOT / "requirements" / "normalized" / safe_id(state.upper())
            out_dir.mkdir(parents=True, exist_ok=True)
            out_file = out_dir / f"{safe_id(req_id)}.json"
            if out_file.exists() and not args.overwrite:
                skipped += 1
                continue
            out_file.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            imported += 1

    print(f"Imported: {imported}")
    print(f"Skipped:  {skipped}")
    print("Next: python tools/build_indexes.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
'''

FILES['tools/build_indexes.py'] = r'''#!/usr/bin/env python3
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set

ROOT = Path(__file__).resolve().parents[1]

FIELD_CANDIDATES = {
    "id": ["id", "ruleId", "ruleID", "rule_id", "code", "name"],
    "coverage": ["coverage", "coverageCode", "coverageName", "coverage_id", "targetCoverage", "coverageCd"]
}


def read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"WARNING: cannot parse {path}: {exc}")
        return None


def first_present(obj: Dict[str, Any], names: Iterable[str]) -> Optional[Any]:
    for name in names:
        if name in obj and obj[name] not in (None, ""):
            return obj[name]
    return None


def normalize_rules(raw: Any) -> List[Dict[str, Any]]:
    if isinstance(raw, list):
        return [x for x in raw if isinstance(x, dict)]
    if isinstance(raw, dict):
        for key in ("rules", "data", "items", "results"):
            if isinstance(raw.get(key), list):
                return [x for x in raw[key] if isinstance(x, dict)]
        return [raw]
    return []


def discover_coverages(obj: Any, depth: int = 0) -> Set[str]:
    if depth > 6:
        return set()
    found: Set[str] = set()
    if isinstance(obj, dict):
        for key, value in obj.items():
            if "coverage" in key.lower() and isinstance(value, (str, int, float)):
                found.add(str(value))
            found.update(discover_coverages(value, depth + 1))
    elif isinstance(obj, list):
        for value in obj:
            found.update(discover_coverages(value, depth + 1))
    return {v for v in found if v and len(v) <= 100}


def ensure_rule_files(state_dir: Path) -> List[Path]:
    rules_dir = state_dir / "rules"
    rules_dir.mkdir(parents=True, exist_ok=True)
    existing = sorted(rules_dir.glob("*.json"))
    if existing:
        return existing

    raw_file = state_dir / "raw" / "ccv-rules.json"
    if not raw_file.exists():
        return []
    rules = normalize_rules(read_json(raw_file))
    created = []
    for i, rule in enumerate(rules, start=1):
        rid = first_present(rule, FIELD_CANDIDATES["id"])
        rid = str(rid) if rid is not None else f"{state_dir.name}_RULE_{i:04d}"
        safe = "".join(c if c.isalnum() or c in "._-" else "-" for c in rid)
        path = rules_dir / f"{safe}.json"
        if path.exists():
            path = rules_dir / f"{safe}-{i:04d}.json"
        path.write_text(json.dumps(rule, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        created.append(path)
    return created


def main() -> int:
    indexes = ROOT / "indexes"
    indexes.mkdir(parents=True, exist_ok=True)

    rules_by_id = {}
    rules_by_state = defaultdict(list)
    rules_by_coverage = defaultdict(lambda: defaultdict(list))

    states_root = ROOT / "states"
    for state_dir in sorted(p for p in states_root.iterdir() if p.is_dir()):
        state = state_dir.name.upper()
        for path in ensure_rule_files(state_dir):
            rule = read_json(path)
            if not isinstance(rule, dict):
                continue
            rid = first_present(rule, FIELD_CANDIDATES["id"])
            rid = str(rid) if rid is not None else path.stem
            rel = str(path.relative_to(ROOT)).replace("\\", "/")
            rules_by_id[rid] = {"state": state, "file": rel}
            rules_by_state[state].append(rid)
            coverages = discover_coverages(rule)
            direct = first_present(rule, FIELD_CANDIDATES["coverage"])
            if direct is not None:
                if isinstance(direct, list):
                    coverages.update(str(x) for x in direct)
                else:
                    coverages.add(str(direct))
            for cov in sorted(coverages):
                rules_by_coverage[state][cov].append(rid)

    req_by_state = defaultdict(list)
    req_by_coverage = defaultdict(lambda: defaultdict(list))
    req_root = ROOT / "requirements" / "normalized"
    if req_root.exists():
        for path in sorted(req_root.rglob("*.json")):
            req = read_json(path)
            if not isinstance(req, dict):
                continue
            rid = str(req.get("id") or path.stem)
            state = str(req.get("state") or "UNKNOWN").upper()
            coverage = str(req.get("coverage") or "UNKNOWN")
            req_by_state[state].append(rid)
            req_by_coverage[state][coverage].append(rid)

    outputs = {
        "rules-by-id.json": rules_by_id,
        "rules-by-state.json": dict(rules_by_state),
        "rules-by-coverage.json": {s: dict(v) for s, v in rules_by_coverage.items()},
        "requirements-by-state.json": dict(req_by_state),
        "requirements-by-coverage.json": {s: dict(v) for s, v in req_by_coverage.items()}
    }

    for name, data in outputs.items():
        (indexes / name).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"Wrote indexes/{name}")

    print("\nIMPORTANT: FIELD_CANDIDATES is generic.")
    print("After seeing the actual internal export, update the exact rule-ID/coverage field names.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
'''

FILES['tools/validate_structure.py'] = r'''#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, List, Set, Tuple

ROOT = Path(__file__).resolve().parents[1]

STANDARD = {
    "==", "===", "!=", "!==", ">", ">=", "<", "<=", "!", "!!",
    "or", "and", "if", "?:", "var", "missing", "missing_some", "in",
    "cat", "substr", "+", "-", "*", "/", "%", "min", "max", "map",
    "filter", "reduce", "all", "none", "some", "merge", "log"
}

XGW = {
    "x-gw-cc-rule-info", "x-gw-cc-rule-success", "x-gw-cc-rule-error",
    "x-gw-default", "x-gw-forbidden", "x-gw-forbidden-property"
}


def custom_names() -> Set[str]:
    path = ROOT / "knowledge" / "json-logic" / "custom-operators.yaml"
    names = set()
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            s = line.strip()
            if s.startswith("- name:"):
                n = s.split(":", 1)[1].strip().strip("'\"")
                if n:
                    names.add(n)
    return names


def walk(obj: Any, loc: str = "$") -> Tuple[Set[str], List[Tuple[str, str]]]:
    ops: Set[str] = set()
    responses: List[Tuple[str, str]] = []
    if isinstance(obj, dict):
        if len(obj) == 1:
            k = next(iter(obj))
            if k in STANDARD or k in custom_names():
                ops.add(k)
        for k, v in obj.items():
            child = f"{loc}.{k}"
            if k in XGW:
                responses.append((child, k))
            co, cr = walk(v, child)
            ops.update(co)
            responses.extend(cr)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            co, cr = walk(v, f"{loc}[{i}]")
            ops.update(co)
            responses.extend(cr)
    return ops, responses


def validate(path: Path) -> int:
    print(f"\n=== {path} ===")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"LIKELY DEFECT: invalid JSON: {exc}")
        return 1
    print("PASS: valid JSON")
    ops, responses = walk(data)
    if ops:
        print("Recognized operators:", ", ".join(sorted(ops)))
    if custom_names():
        print("Documented custom operators:", ", ".join(sorted(custom_names())))
    if responses:
        print("CCV response properties:")
        for loc, name in responses:
            print(f"  - {name} at {loc}")
    print("NOTE: semantic correctness still requires requirement comparison + runtime tests.")
    return 0


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("file", nargs="?", type=Path)
    p.add_argument("--state")
    args = p.parse_args()
    files: List[Path] = []
    if args.file:
        files = [args.file]
    elif args.state:
        rules_dir = ROOT / "states" / args.state.upper() / "rules"
        files = sorted(rules_dir.glob("*.json"))
        if not files:
            raw = ROOT / "states" / args.state.upper() / "raw" / "ccv-rules.json"
            if raw.exists():
                files = [raw]
    else:
        p.error("provide a file or --state")
    if not files:
        print("No rule files found.")
        return 2
    failures = sum(validate(f) for f in files)
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
'''

FILES['tools/compare_states.py'] = r'''#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def load(name: str):
    path = ROOT / "indexes" / name
    if not path.exists():
        raise SystemExit(f"Missing {path}. Run python tools/build_indexes.py first.")
    return json.loads(path.read_text(encoding="utf-8"))


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("state_a")
    p.add_argument("state_b")
    args = p.parse_args()
    a, b = args.state_a.upper(), args.state_b.upper()
    rs = load("rules-by-state.json")
    rc = load("rules-by-coverage.json")
    qs = load("requirements-by-state.json")
    qc = load("requirements-by-coverage.json")
    print(f"{a} vs {b}")
    print("=" * 60)
    print(f"Rules:        {a}={len(rs.get(a, []))}  {b}={len(rs.get(b, []))}")
    print(f"Requirements: {a}={len(qs.get(a, []))}  {b}={len(qs.get(b, []))}")
    covs = sorted(set(rc.get(a, {})) | set(rc.get(b, {})) | set(qc.get(a, {})) | set(qc.get(b, {})))
    for c in covs:
        print(f"{c:25} rules {a}:{len(rc.get(a, {}).get(c, [])):3} {b}:{len(rc.get(b, {}).get(c, [])):3} | reqs {a}:{len(qc.get(a, {}).get(c, [])):3} {b}:{len(qc.get(b, {}).get(c, [])):3}")
    print("\nStructural comparison only; use CCV Expert for semantic comparison.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
'''

FILES['tools/README.md'] = r'''
# Tools

- `import_requirements.py`: import existing work Excel requirements into small JSON files.
- `build_indexes.py`: split/index rule exports and requirements for token-efficient retrieval.
- `validate_structure.py`: conservative deterministic JSON/operator/x-gw checks.
- `compare_states.py`: structural state comparison.

The actual system-table rule shape is internal. After inspecting it, update
`FIELD_CANDIDATES` in `build_indexes.py` so indexing matches the exact fields.
'''

FILES['indexes/.gitkeep'] = ''
FILES['reports/.gitkeep'] = ''

STATES = {
    'WI': ('Wisconsin', 'RELEASED', True),
    'IL': ('Illinois', 'DEVELOPMENT', False),
    'OH': ('Ohio', 'DEVELOPMENT', False),
    'AZ': ('Arizona', 'DEVELOPMENT', False),
    'UT': ('Utah', 'REFINEMENT', False),
    'IN': ('Indiana', 'REFINEMENT', False),
}

for code, (name, status, reference) in STATES.items():
    FILES[f'states/{code}/metadata.json'] = json.dumps({
        'state': code,
        'name': name,
        'status': status,
        'releasedReference': reference,
        'notes': ''
    }, indent=2) + '\n'
    for sub in ('raw', 'rules', 'coverages', 'ui-rules', 'tests', 'indexes'):
        FILES[f'states/{code}/{sub}/.gitkeep'] = ''


def normalize(content: str) -> str:
    content = textwrap.dedent(content).lstrip('\n')
    if content and not content.endswith('\n'):
        content += '\n'
    return content


def main() -> int:
    p = argparse.ArgumentParser(description='Create the generic CCV Engineering Assistant workspace.')
    p.add_argument('--root', default='ccv-assistant', help='Output directory (default: ccv-assistant)')
    p.add_argument('--force', action='store_true', help='Overwrite existing generated files')
    args = p.parse_args()

    root = Path(args.root).resolve()
    print('CCV Engineering Assistant Bootstrap')
    print('=' * 60)
    print(f'Target: {root}')
    print()

    written = skipped = 0
    for rel, content in sorted(FILES.items()):
        path = root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists() and not args.force:
            print(f'SKIP    {rel}')
            skipped += 1
            continue
        path.write_text(normalize(content), encoding='utf-8', newline='\n')
        print(f'CREATE  {rel}')
        written += 1

    print('\n' + '=' * 60)
    print(f'Created/updated: {written}')
    print(f'Skipped:         {skipped}')
    print('\nNEXT ACTIONS')
    print('1. Open the generated ccv-assistant folder in VS Code.')
    print('2. Add states/<STATE>/raw/ccv-rules.json from the internal Postman/system-table export.')
    print('3. Add states/<STATE>/raw/coverages.json from internal Guidewire/coverage data.')
    print('4. Complete knowledge/json-logic/custom-operators.yaml.')
    print('5. Complete knowledge/ccv-response/response-contract.yaml.')
    print('6. Verify/add UI-coded CCV behavior in knowledge/implementation-patterns/ui-coded-rules.md.')
    print('7. Inspect internal requirements workbook:')
    print('     python tools/import_requirements.py <requirements.xlsx> --inspect')
    print('8. Import requirements using exact workbook headers.')
    print('9. Run: python tools/build_indexes.py')
    print('10. Validate representative Wisconsin rules before development-state validation.')
    print('\nNo internal/company-specific data is included in this bootstrap.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
